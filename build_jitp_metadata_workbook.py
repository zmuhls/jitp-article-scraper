#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


DEFAULT_OUTPUT_DIR = Path.cwd() / "outputs" / "jitp_manifold_metadata"
EXCEL_TEXT_LIMIT = 32767


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def compact(value):
    return re.sub(r"\s+", " ", as_text(value)).strip()


def trim_cell(value, limit=EXCEL_TEXT_LIMIT):
    text = compact(value)
    if len(text) <= limit:
        return text
    return text[: limit - 30].rstrip() + " ... [truncated for Excel]"


def unique(values):
    seen = []
    keys = set()
    for value in values:
        text = compact(value)
        if text and text not in keys:
            keys.add(text)
            seen.append(text)
    return seen


def join_unique(values, separator=" | "):
    return separator.join(unique(values))


def normalize_name(name):
    return compact(name).lower().replace("'", "").replace("’", "")


def date_only(value):
    text = compact(value)
    if not text:
        return ""
    match = re.match(r"^(\d{4}-\d{2}-\d{2})", text)
    return match.group(1) if match else text


def parse_json_object(value):
    text = compact(value)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def best_publication_date(row):
    metadata = parse_json_object(row.get("text_metadata_json"))
    candidates = [
        ("text_publication_date", row.get("text_publication_date")),
        ("issue_publication_date", row.get("issue_publication_date")),
        ("text_metadata.originalPublicationDate", metadata.get("originalPublicationDate")),
        ("text_created_at", row.get("text_created_at")),
        ("text_updated_at", row.get("text_updated_at")),
    ]
    for source, value in candidates:
        if compact(value):
            return date_only(value), source
    return "", ""


def note_for(row):
    title = compact(row.get("title")).lower()
    if row.get("deleted_slug_flag") == "yes":
        return "Deleted/404 text entry exposed by Manifold API."
    if any(label in title for label in ("guidelines", "call for submissions", "about the journal")):
        return "Likely journal admin/guideline text included in a Manifold issue project."
    if row.get("bio_status") != "found":
        return "No author bio visible in rendered article page."
    if row.get("article_html_status") != "parsed":
        return "Article page did not parse cleanly."
    return ""


def split_heading_text(value):
    return [re.sub(r"\[\d+\]", "", part).strip() for part in compact(value).split("|") if part.strip()]


def clean_affiliation(value):
    text = compact(value)
    text = re.sub(
        r"\s+(As we|As instructors|This essay|This article|This paper|In this|With the|Although existing|Though use|Open Educational|Pandemic teaching)\b.*$",
        "",
        text,
        flags=re.I,
    )
    text = re.sub(r"\s+##.*$", "", text)
    return text.strip()


def parse_author_list_segment(value, role="Author", source="author_names"):
    text = compact(value)
    text = re.sub(r"^by\s+", "", text, flags=re.I)
    if not text:
        return []
    text = re.sub(r"\s*&\s*", " and ", text)
    text = re.sub(r",\s*(Jr\.?|Sr\.?|II|III|IV)\b", r"§SUFFIX§ \1", text, flags=re.I)
    text = re.sub(r"\s+and\s+", ", ", text, flags=re.I)
    records = []
    for name in re.split(r"\s*,\s*", text):
        cleaned = name.replace("§SUFFIX§", ",").strip()
        if cleaned:
            records.append({"name": cleaned, "role": role, "affiliation": "", "source": source})
    return records


def parse_author_string(value, source="author_names"):
    text = compact(value)
    if not text:
        return []
    records = []
    for segment in re.split(r"\s*\|\s*", text):
        match = re.match(r"^(.*?)(?:,?\s+in conversation with\s+)(.+)$", segment, flags=re.I)
        if match:
            records.extend(parse_author_list_segment(match.group(1), "Author", source))
            records.extend(parse_author_list_segment(match.group(2), "Conversation participant", source))
        else:
            records.extend(parse_author_list_segment(segment, "Author", source))
    return [record for record in records if record["name"]]


def parse_byline_line(line):
    parts = [part.strip() for part in re.sub(r"\[\d+\]", "", compact(line)).split(",") if part.strip()]
    if not parts:
        return "", ""
    if len(parts) > 1 and re.match(r"^(Jr\.?|Sr\.?|II|III|IV)$", parts[1], flags=re.I):
        return f"{parts[0]}, {parts[1]}", clean_affiliation(", ".join(parts[2:]))
    return parts[0], clean_affiliation(", ".join(parts[1:]))


SECTION_STOP_RE = re.compile(
    r"^(abstract|introduction|background|method|methods|results|discussion|conclusion|notes?|references|bibliography|works cited|acknowledg(?:e)?ments?|appendix|about the author|about the authors|about the editors|mission statement|review policy|masthead|how well|teaching with|the openlab|course overview|literature review)$",
    flags=re.I,
)


def bio_author_names(bio):
    text = compact(bio)
    if not text:
        return []
    pattern = re.compile(
        r"(?:^|[.!?]\s+)([A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+(?:\s+(?:[A-Z]\.|[A-Z][A-Za-zÀ-ÖØ-öø-ÿ'’.-]+|de|del|van|von)){1,6})\s+(?:is|are|was|were|earned|received|leads|serves|completed|holds|has|works|teaches|studies|graduated|coordinates|directs|develops|creates|researches|uses|conducted|sits|currently)\b"
    )
    return unique(match.group(1) for match in pattern.finditer(text))


def filled_bio(row):
    if compact(row.get("author_bio")):
        return row.get("author_bio"), "scraped_bio" if row.get("bio_status") == "found" else "raw_bio"
    text = compact(row.get("normalized_full_text"))
    match = re.search(r"\s##\s+About the Authors?\b\s+([\s\S]*?)(?=\s##\s+[A-Z0-9]|\sThis entry is licensed|\sBack to top|$)", text)
    if match:
        return compact(match.group(1)), "normalized_text_about_section"
    return "", ""


def filled_abstract(row):
    if compact(row.get("abstract")):
        return row.get("abstract"), "scraped_abstract" if row.get("abstract_status") == "found" else "raw_abstract"
    text = compact(row.get("normalized_full_text"))
    match = re.search(r"\s##\s+Abstract\b\s+([\s\S]*?)(?=\s##\s+[A-Z0-9]|\sThis entry is licensed|\sBack to top|$)", text)
    if match:
        return compact(match.group(1)), "normalized_text_abstract_section"
    return "", ""


def heading_byline_entries(row):
    headings = split_heading_text(row.get("author_byline_lines") or row.get("section_headings") or row.get("toc_headings"))
    if not headings:
        return []
    title = compact(row.get("title")).lower()
    try:
        start = [heading.lower() for heading in headings].index(title) + 1
    except ValueError:
        start = 1
    names = bio_author_names(row.get("author_bio"))
    bylines = []
    for heading in headings[start:]:
        if SECTION_STOP_RE.match(heading):
            break
        if len(heading) > 220:
            break
        if "," in heading or any(name in heading for name in names):
            bylines.append(heading)
    source = "rendered_byline" if row.get("author_byline_lines") else "inferred_heading"
    return [{"line": line, "source": source} for line in bylines]


def text_prefix_byline_entries(row):
    text = compact(row.get("normalized_full_text"))
    title = compact(row.get("title"))
    if not text:
        return []
    rest = re.sub(r"^#+\s*", "", text)
    if title and rest.lower().startswith(title.lower()):
        rest = rest[len(title) :].strip()
    if rest.startswith("###"):
        before_body = re.split(r"\s+##\s+", rest, maxsplit=1)[0]
        lines = [
            re.sub(r"^###\s*", "", line).strip()
            for line in re.split(r"\s+###\s+", before_body)
        ]
        return [{"line": line, "source": "inferred_text_heading"} for line in lines if line and not SECTION_STOP_RE.match(line)]
    abstract_index = re.search(r"\s+##\s+Abstract\b", rest, flags=re.I)
    if abstract_index:
        candidate = rest[: abstract_index.start()].strip()
        if candidate and len(candidate) < 1200:
            return [{"line": candidate, "source": "inferred_text_prefix"}]
    return []


def byline_entries_for_article(row):
    rendered = [{"line": line, "source": "rendered_byline"} for line in split_heading_text(row.get("author_byline_lines"))]
    if rendered:
        return rendered
    heading = heading_byline_entries(row)
    if heading:
        return heading
    return text_prefix_byline_entries(row)


def split_byline_blob(blob, bio_names=None):
    text = re.sub(r"\[\d+\]", "", compact(blob))
    bio_names = bio_names or []
    if not text:
        return []
    anchors = sorted(
        ((name, text.find(name)) for name in bio_names if text.find(name) >= 0),
        key=lambda item: item[1],
    )
    if len(anchors) > 1:
        return [text[index : anchors[pos + 1][1] if pos + 1 < len(anchors) else len(text)].strip() for pos, (_, index) in enumerate(anchors)]
    return split_heading_text(text) if " | " in text else [text]


def records_from_byline_entries(entries, bio_names):
    records = []
    for entry in entries:
        for segment in split_byline_blob(entry["line"], bio_names):
            name, affiliation = parse_byline_line(segment)
            parsed = parse_author_string(name, entry["source"])
            if not parsed and name:
                parsed = [{"name": name, "role": "Author", "affiliation": "", "source": entry["source"]}]
            for record in parsed:
                record["affiliation"] = affiliation or record.get("affiliation", "")
                record["source"] = entry["source"]
                records.append(record)
    return records


def apply_row_affiliations(records, row):
    affiliations = split_heading_text(row.get("author_affiliations"))
    if not affiliations:
        return records
    out = []
    for index, record in enumerate(records):
        new_record = dict(record)
        if not new_record.get("affiliation"):
            new_record["affiliation"] = affiliations[index] if index < len(affiliations) else (affiliations[0] if len(affiliations) == 1 else "")
        out.append(new_record)
    return out


def bios_by_author(row, records):
    bio, _ = filled_bio(row)
    if not bio or not records:
        return {}
    if len(records) == 1:
        index = bio.find(records[0]["name"])
        return {records[0]["split_key"]: bio[index:].strip() if index > 0 else bio}
    anchors = sorted(
        (
            {"key": record["split_key"], "name": record["name"], "index": bio.find(record["name"])}
            for record in records
            if bio.find(record["name"]) >= 0
        ),
        key=lambda item: item["index"],
    )
    if not anchors:
        return {}
    out = {}
    for index, anchor in enumerate(anchors):
        end = anchors[index + 1]["index"] if index + 1 < len(anchors) else len(bio)
        out[anchor["key"]] = bio[anchor["index"] : end].strip()
    return out


def author_records_for_article(row):
    bio, bio_source = filled_bio(row)
    bio_names = bio_author_names(bio)
    rendered = [{"line": line, "source": "rendered_byline"} for line in split_heading_text(row.get("author_byline_lines"))]
    if rendered:
        records = records_from_byline_entries(rendered, bio_names)
    elif compact(row.get("author_names")):
        records = apply_row_affiliations(parse_author_string(row.get("author_names"), "api_description"), row)
    else:
        records = records_from_byline_entries(byline_entries_for_article(row), bio_names)

    keyed = [{**record, "split_key": normalize_name(record.get("name"))} for record in records]
    author_bios = bios_by_author(row, keyed)
    abstract, abstract_source = filled_abstract(row)
    date, date_source = best_publication_date(row)
    out = []
    for index, record in enumerate(keyed, start=1):
        out.append(
            {
                **record,
                "author_position": index,
                "author_count_for_article": len(keyed),
                "record_id": row.get("record_id"),
                "text_id": row.get("text_id"),
                "project_id": row.get("project_id"),
                "issue_number": row.get("issue_number"),
                "issue_type": row.get("issue_type"),
                "short_form_section": row.get("short_form_section"),
                "text_category": row.get("text_category"),
                "text_position": row.get("text_position"),
                "title": row.get("title"),
                "subtitle": row.get("subtitle"),
                "article_url": row.get("article_url"),
                "project_url": row.get("project_url"),
                "text_publication_date": row.get("text_publication_date"),
                "issue_publication_date": row.get("issue_publication_date"),
                "best_publication_date": date,
                "best_publication_date_source": date_source,
                "article_author_cell": row.get("author_names"),
                "article_author_cell_filled": row.get("author_names") or " | ".join(item["name"] for item in keyed),
                "article_bio_cell": row.get("author_bio"),
                "article_bio_cell_filled": bio,
                "article_bio_source": bio_source,
                "author_bio": author_bios.get(record["split_key"], ""),
                "abstract": abstract,
                "abstract_source": abstract_source,
                "keywords": row.get("keywords"),
                "bio_status": row.get("bio_status"),
                "abstract_status": row.get("abstract_status"),
                "keywords_status": row.get("keywords_status"),
                "license": row.get("license"),
                "record_note": note_for(row),
                "existing_result_used": row.get("existing_result_used"),
            }
        )
    return out


def issue_sort_key(value):
    text = as_text(value)
    return (0, int(text)) if text.isdigit() else (1, text)


def sort_article_key(row):
    issue_group, issue_value = issue_sort_key(row.get("issue_number"))
    return (issue_group, issue_value, int(row.get("text_position") or 999), compact(row.get("title")))


def make_publication_row(row, records_by_article, stats_by_number, stats_by_section):
    authors = records_by_article.get(row.get("record_id"), [])
    date, date_source = best_publication_date(row)
    abstract, abstract_source = filled_abstract(row)
    bio, bio_source = filled_bio(row)
    stats = stats_by_section.get(row.get("issue_number")) if row.get("issue_type") == "Short Form Section" else stats_by_number.get(row.get("issue_number"))
    out = {
        "record_id": row.get("record_id"),
        "issue_number": row.get("issue_number"),
        "issue_title": row.get("issue_title"),
        "issue_subtitle": row.get("issue_subtitle"),
        "issue_editors": row.get("issue_editors"),
        "issue_type": row.get("issue_type"),
        "short_form_section": row.get("short_form_section") or ("" if as_text(row.get("issue_number")).isdigit() else row.get("issue_number")),
        "text_category": row.get("text_category"),
        "text_position": row.get("text_position"),
        "title": row.get("title"),
        "subtitle": row.get("subtitle"),
        "best_publication_date": date,
        "best_publication_date_source": date_source,
        "text_publication_date": date_only(row.get("text_publication_date")),
        "issue_publication_date": date_only(row.get("issue_publication_date")),
        "author_count": len(authors),
        "authors_combined": join_unique(record.get("name") for record in authors),
        "author_affiliations_combined": join_unique(": ".join(part for part in (record.get("name"), record.get("affiliation")) if part) for record in authors),
        "raw_author_cell": row.get("author_names"),
        "filled_author_cell": row.get("author_names") or " | ".join(record.get("name", "") for record in authors),
        "author_split_sources": join_unique((record.get("source") for record in authors), ", "),
        "abstract": abstract,
        "abstract_source": abstract_source,
        "keywords": row.get("keywords"),
        "keywords_status": row.get("keywords_status"),
        "article_level_bio": bio,
        "article_bio_source": bio_source,
        "bio_status": row.get("bio_status"),
        "normalized_main_text": row.get("normalized_main_text"),
        "notes_text": row.get("notes_text"),
        "references_text": row.get("references_text"),
        "acknowledgments_text": row.get("acknowledgments_text"),
        "appendix_text": row.get("appendix_text"),
        "normalized_main_word_count": row.get("normalized_main_word_count"),
        "normalized_full_word_count": row.get("normalized_full_word_count"),
        "paragraph_count": row.get("paragraph_count"),
        "heading_count": row.get("heading_count"),
        "section_headings": row.get("section_headings"),
        "toc_headings": row.get("toc_headings"),
        "license": row.get("license"),
        "article_url": row.get("article_url"),
        "project_url": row.get("project_url"),
        "text_id": row.get("text_id"),
        "project_id": row.get("project_id"),
        "text_slug": row.get("text_slug"),
        "text_doi": row.get("text_doi"),
        "text_unique_identifier": row.get("text_unique_identifier"),
        "project_text_records_found": stats.get("text_records_found", "") if stats else "",
        "project_author_instances_found": stats.get("author_instances_found", "") if stats else "",
        "project_unique_authors_found": stats.get("unique_authors_found", "") if stats else "",
        "record_note": note_for(row),
        "existing_result_used": row.get("existing_result_used"),
    }
    for index, record in enumerate(authors, start=1):
        out[f"author_{index}_name"] = record.get("name")
        out[f"author_{index}_role"] = record.get("role")
        out[f"author_{index}_affiliation"] = record.get("affiliation")
        out[f"author_{index}_bio"] = record.get("author_bio")
    return out


def add_table(ws, columns, rows, table_name):
    header_fill = PatternFill("solid", fgColor="1F4E5F")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, (_, label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(columns, start=1):
            value = key(row) if callable(key) else row.get(key, "")
            ws.cell(row=row_idx, column=col_idx, value=trim_cell(value))
        ws.row_dimensions[row_idx].height = 66
    end_row = max(1, len(rows) + 1)
    end_col = get_column_letter(len(columns))
    table = Table(displayName=table_name, ref=f"A1:{end_col}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "A2"


def set_widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(input_json, output_xlsx):
    raw = json.loads(Path(input_json).read_text(encoding="utf-8"))
    articles = raw["articles"]
    issues = raw["issues"]
    summary = raw.get("summary", {})
    author_records = [record for row in articles for record in author_records_for_article(row)]

    records_by_article = defaultdict(list)
    for record in author_records:
        records_by_article[record["record_id"]].append(record)
    max_author_slots = max([1] + [len(records) for records in records_by_article.values()])

    records_by_author = defaultdict(list)
    for record in author_records:
        if record.get("split_key"):
            records_by_author[record["split_key"]].append(record)

    author_rows = []
    for key, records in records_by_author.items():
        sorted_records = sorted(records, key=lambda record: (record.get("best_publication_date") or "", sort_article_key(record)))
        dates = sorted(unique(record.get("best_publication_date") or record.get("text_publication_date") or record.get("issue_publication_date") for record in records))
        author_rows.append(
            {
                "author_name": sorted_records[0].get("name"),
                "author_name_normalized": key,
                "publication_count": len(set(record.get("record_id") for record in records)),
                "author_record_count": len(records),
                "first_publication_date": dates[0] if dates else "",
                "latest_publication_date": dates[-1] if dates else "",
                "numbered_issue_records": sum(1 for record in records if record.get("issue_type") == "Numbered Issue"),
                "short_form_records": sum(1 for record in records if record.get("issue_type") == "Short Form Section"),
                "issues_sections": join_unique((record.get("issue_number") for record in records), ", "),
                "roles": join_unique((record.get("role") for record in records), ", "),
                "affiliations": join_unique(record.get("affiliation") for record in records),
                "bio_statuses": join_unique((record.get("bio_status") for record in records), ", "),
                "representative_bio": next((record.get("author_bio") for record in sorted_records if compact(record.get("author_bio"))), ""),
                "article_titles": join_unique(record.get("title") for record in records),
                "article_urls": join_unique(record.get("article_url") for record in records),
                "source_author_cells": join_unique(record.get("article_author_cell_filled") for record in records),
                "split_sources": join_unique((record.get("source") for record in records), ", "),
                "record_notes": join_unique(record.get("record_note") for record in records),
            }
        )
    author_rows.sort(key=lambda row: compact(row.get("author_name")).lower())

    def stats_for_issue(issue):
        issue_articles = [row for row in articles if row.get("issue_number") == issue.get("issue_number")]
        issue_authors = [record for record in author_records if record.get("issue_number") == issue.get("issue_number")]
        return {
            **issue,
            "text_records_found": len(issue_articles),
            "author_instances_found": len(issue_authors),
            "unique_authors_found": len(set(record.get("split_key") for record in issue_authors if record.get("split_key"))),
        }

    issue_summary_rows = [stats_for_issue(issue) for issue in issues if as_text(issue.get("issue_number")).isdigit()]
    short_section_rows = [stats_for_issue(issue) for issue in issues if issue.get("issue_type") == "Short Form Section"]
    stats_by_number = {row.get("issue_number"): row for row in issue_summary_rows}
    stats_by_section = {row.get("issue_number"): row for row in short_section_rows}

    issue_rows = [
        make_publication_row(row, records_by_article, stats_by_number, stats_by_section)
        for row in sorted((row for row in articles if as_text(row.get("issue_number")).isdigit()), key=sort_article_key)
    ]
    short_rows = [
        make_publication_row(row, records_by_article, stats_by_number, stats_by_section)
        for row in sorted((row for row in articles if row.get("issue_type") == "Short Form Section"), key=sort_article_key)
    ]

    wb = Workbook()
    wb.remove(wb.active)
    ws_authors = wb.create_sheet("Author Metadata")
    ws_author_pub = wb.create_sheet("Author-Pub Records")
    ws_issues = wb.create_sheet("Issues")
    ws_shorts = wb.create_sheet("Shorts")
    ws_field = wb.create_sheet("Field Notes")

    author_columns = [
        ("author_name", "Author Name"),
        ("author_name_normalized", "Normalized Name"),
        ("publication_count", "Publication Count"),
        ("author_record_count", "Author-Record Count"),
        ("first_publication_date", "First Publication Date"),
        ("latest_publication_date", "Latest Publication Date"),
        ("numbered_issue_records", "Numbered Issue Records"),
        ("short_form_records", "Short-Form Records"),
        ("issues_sections", "Issues / Sections"),
        ("roles", "Roles"),
        ("affiliations", "Affiliations"),
        ("bio_statuses", "Bio Statuses"),
        ("representative_bio", "Representative Bio"),
        ("article_titles", "Article Titles"),
        ("article_urls", "Article URLs"),
        ("source_author_cells", "Source Author Cells"),
        ("split_sources", "Split Sources"),
        ("record_notes", "Record Notes"),
    ]
    add_table(ws_authors, author_columns, author_rows, "AuthorMetadataTable")
    set_widths(ws_authors, [28, 24, 14, 14, 16, 16, 16, 16, 24, 18, 38, 18, 70, 70, 48, 54, 18, 40])

    author_pub_columns = [
        ("record_id", "Record ID"),
        ("author_position", "Author Position"),
        ("author_count_for_article", "Article Author Count"),
        ("name", "Author Name"),
        ("role", "Role"),
        ("affiliation", "Affiliation"),
        ("source", "Split Source"),
        ("article_author_cell", "Original Author Cell"),
        ("article_author_cell_filled", "Filled Author Cell"),
        ("issue_number", "Issue / Section"),
        ("issue_type", "Issue Type"),
        ("short_form_section", "Short Form Section"),
        ("text_category", "Text Category"),
        ("text_position", "Text Position"),
        ("title", "Article Title"),
        ("subtitle", "Subtitle"),
        ("best_publication_date", "Best Publication Date"),
        ("best_publication_date_source", "Date Source"),
        ("abstract", "Abstract"),
        ("abstract_source", "Abstract Source"),
        ("keywords", "Keywords"),
        ("bio_status", "Bio Status"),
        ("author_bio", "Author Bio"),
        ("article_bio_cell_filled", "Article-Level Bio"),
        ("article_bio_source", "Bio Source"),
        ("license", "License"),
        ("record_note", "Record Note"),
        ("existing_result_used", "Existing Result Fallback Fields"),
        ("article_url", "Article URL"),
        ("project_url", "Project URL"),
        ("text_id", "Text ID"),
        ("project_id", "Project ID"),
    ]
    add_table(ws_author_pub, author_pub_columns, author_records, "AuthorPubRecordsTable")
    set_widths(ws_author_pub, [28, 12, 12, 26, 18, 30, 18, 40, 40, 12, 16, 16, 18, 10, 42, 26, 14, 18, 60, 18, 30, 14, 70, 70, 18, 36, 34, 34, 42, 36, 28, 28])

    author_slot_columns = []
    for slot in range(1, max_author_slots + 1):
        author_slot_columns.extend(
            [
                (f"author_{slot}_name", f"Author {slot} Name"),
                (f"author_{slot}_role", f"Author {slot} Role"),
                (f"author_{slot}_affiliation", f"Author {slot} Affiliation"),
                (f"author_{slot}_bio", f"Author {slot} Bio"),
            ]
        )

    publication_columns = [
        ("record_id", "Record ID"),
        ("issue_number", "Issue / Section"),
        ("issue_title", "Issue / Project Title"),
        ("issue_subtitle", "Issue Subtitle"),
        ("issue_editors", "Issue Editors"),
        ("issue_type", "Issue Type"),
        ("short_form_section", "Short Form Section"),
        ("text_category", "Text Category"),
        ("text_position", "Text Position"),
        ("title", "Article Title"),
        ("subtitle", "Subtitle"),
        ("best_publication_date", "Best Publication Date"),
        ("best_publication_date_source", "Date Source"),
        ("author_count", "Author Count"),
        ("authors_combined", "Authors Combined"),
        ("author_affiliations_combined", "Author Affiliations Combined"),
        ("raw_author_cell", "Raw Author Cell"),
        ("filled_author_cell", "Filled Author Cell"),
        ("author_split_sources", "Author Split Sources"),
        *author_slot_columns,
        ("abstract", "Abstract"),
        ("abstract_source", "Abstract Source"),
        ("keywords", "Keywords"),
        ("article_level_bio", "Article-Level Bio"),
        ("article_bio_source", "Bio Source"),
        ("normalized_main_text", "Normalized Main Text"),
        ("notes_text", "Notes Text"),
        ("references_text", "References Text"),
        ("acknowledgments_text", "Acknowledgments Text"),
        ("appendix_text", "Appendix Text"),
        ("normalized_main_word_count", "Main Word Count"),
        ("normalized_full_word_count", "Full Word Count"),
        ("paragraph_count", "Paragraph Count"),
        ("heading_count", "Heading Count"),
        ("section_headings", "Section Headings"),
        ("toc_headings", "TOC Headings"),
        ("license", "License"),
        ("article_url", "Article URL"),
        ("project_url", "Project URL"),
        ("text_id", "Text ID"),
        ("project_id", "Project ID"),
        ("text_slug", "Text Slug"),
        ("text_doi", "Text DOI"),
        ("text_unique_identifier", "Text Unique Identifier"),
        ("project_text_records_found", "Project Text Records Found"),
        ("project_author_instances_found", "Project Author Instances Found"),
        ("project_unique_authors_found", "Project Unique Authors Found"),
        ("record_note", "Record Note"),
        ("existing_result_used", "Existing Result Fallback Fields"),
    ]
    add_table(ws_issues, publication_columns, issue_rows, "IssuesTable")
    add_table(ws_shorts, publication_columns, short_rows, "ShortsTable")
    widths = [28, 14, 30, 28, 28, 16, 18, 16, 10, 42, 24, 14, 18, 12, 40, 42, 40, 40, 20]
    widths.extend([24, 14, 30, 70] * max_author_slots)
    widths.extend([60, 18, 30, 70, 18, 80, 55, 65, 45, 45, 14, 14, 12, 12, 55, 55, 36, 42, 36, 28, 28, 28, 24, 24, 12, 14, 14, 36, 34])
    set_widths(ws_issues, widths)
    set_widths(ws_shorts, widths)

    short_counts = " | ".join(f"{row.get('section_name') or row.get('issue_number')}: {row.get('text_records_found')}" for row in short_section_rows)
    field_rows = [
        {"field": "Workbook Scope", "note": "Focused JITP Manifold metadata workbook with author-level metadata, author-publication records, article-level numbered issue records, article-level short-form records, and field notes."},
        {"field": "Source Base", "note": summary.get("source_base_url", "")},
        {"field": "Scraped At UTC", "note": summary.get("scraped_at_utc", "")},
        {"field": "Text Records", "note": summary.get("article_record_count", len(articles))},
        {"field": "Author Instances Found", "note": len(author_records)},
        {"field": "Unique Author Names Found", "note": len(records_by_author)},
        {"field": "Numbered Issues", "note": len(issue_summary_rows)},
        {"field": "Numbered Issue Publications", "note": len(issue_rows)},
        {"field": "Short-Form Sections", "note": len(short_section_rows)},
        {"field": "Short-Form Publications", "note": len(short_rows)},
        {"field": "Short-Form Publication Counts", "note": short_counts},
        {"field": "Author Metadata", "note": "One row per normalized author name."},
        {"field": "Author-Pub Records", "note": "One row per parsed author/article relationship."},
        {"field": "Issues", "note": "One row per numbered-issue article, with wide author slots."},
        {"field": "Shorts", "note": "One row per short-form article, with wide author slots."},
        {"field": "QA Notes", "note": "Missing author bios, admin/guideline texts, deleted/404 slugs, and parse issues are surfaced in Record Note columns."},
    ]
    add_table(ws_field, [("field", "Field"), ("note", "Note")], field_rows, "FieldNotesTable")
    set_widths(ws_field, [28, 100])

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)
    print(output_xlsx)


def parse_args():
    parser = argparse.ArgumentParser(description="Build the JITP metadata workbook from scraper JSON.")
    parser.add_argument("--input-json", default=str(DEFAULT_OUTPUT_DIR / "jitp_manifold_metadata.json"))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--output-xlsx", default="")
    return parser.parse_args()


def main():
    args = parse_args()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_xlsx = Path(args.output_xlsx).expanduser().resolve() if args.output_xlsx else output_dir / "jitp_manifold_article_metadata.xlsx"
    build_workbook(Path(args.input_json).expanduser().resolve(), output_xlsx)


if __name__ == "__main__":
    main()
